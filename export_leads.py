"""Scores, deduplicates and exports leads to a styled Excel file.

Sheets:
  Event Leads          upcoming events with a confirmed date
  Date Not Confirmed   events where the date could not be confirmed
  Company Leads        agencies, production companies, venues, organisers
  Summary              headline numbers
"""

import os
import re
import json
from datetime import date
from difflib import SequenceMatcher
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

IN_PATH = os.path.join("data", "extracted.json")
OUT_DIR = "output"
TODAY = date.today()
MIN_DAYS_AWAY = 5

CORE_STATES = {"lagos", "ogun", "oyo", "osun", "kwara"}
CORE_CITIES = {"lagos", "ibadan", "abeokuta", "osogbo", "ilorin", "ikeja",
               "lekki", "victoria island", "ota", "ijebu", "sagamu",
               "ile-ife", "ilisan", "magboro"}
GENERIC_EMAILS = ("info@", "hello@", "contact@", "admin@", "enquiries@",
                  "enquiry@", "support@", "hi@", "mail@", "office@")
FREE_MAIL = {"gmail.com", "yahoo.com", "yahoo.co.uk", "hotmail.com",
             "outlook.com", "icloud.com", "ymail.com"}

# Emails from these domains belong to the website, not the organiser
BAD_EMAIL_DOMAINS = {
    "luma.com", "legit.ng", "corp.legit.ng", "eventbrite.com", "tix.africa",
    "von.gov.ng", "punchng.com", "vanguardngr.com", "guardian.ng",
    "thenationonlineng.net", "tribuneonlineng.com", "channelstv.com",
    "businessday.ng", "bellanaija.com", "pmnewsnigeria.com",
    "leadership.ng", "dailypost.ng", "thisdaylive.com", "lagosarenaguide.com",
}

VENUE_NAME_WORDS = ("centre", "center", "hall", "hotel", "resort",
                    "palace", "towers", "arena")
EXCLUDE_WORDS = ("university", "school", "college", "church", "ministries",
                 "bank", "holdings", "government", "political", "foundation",
                 "corps", "police", "institute", "convention", "religious")

STOP_WORDS = {"the", "2026", "2027", "edition", "annual", "and"}
MAIN_FIELDS = ["name", "start_date", "venue", "city", "organiser", "website"]


# ---------- small helpers ----------

def text(value):
    if value is None:
        return ""
    return str(value).strip()


def parse_date(value):
    try:
        return date.fromisoformat(text(value)[:10])
    except ValueError:
        return None


def domain(value):
    value = text(value)
    if not value:
        return ""
    if "://" not in value:
        value = "http://" + value
    return urlparse(value).netloc.replace("www.", "").lower()


def email_domain(email):
    return email.split("@")[-1].lower() if "@" in email else ""


def domains_match(a, b):
    return bool(a and b and (a == b or a.endswith("." + b) or b.endswith("." + a)))


def pick_email(row):
    candidates = [text(row.get("email")).lower()] + list(row.get("page_emails") or [])
    for email in candidates:
        if "@" not in email:
            continue
        if email_domain(email) in BAD_EMAIL_DOMAINS:
            continue
        return email
    return ""


def pick_phone(row):
    raw = text(row.get("phone"))
    digits = re.sub(r"\D", "", raw)
    if digits.startswith("234"):
        digits = "0" + digits[3:]
    if len(digits) == 11:
        return digits
    found = row.get("page_phones") or []
    return found[0] if found else raw


def needs_verify(row):
    """Flags emails that may belong to a listing site, not the organiser."""
    email = row["_email"]
    if not email:
        return ""
    ed = email_domain(email)
    if ed in FREE_MAIL:
        return ""
    if domains_match(ed, domain(row.get("website"))):
        return ""
    return "Verify"


def in_core_area(row):
    city = text(row.get("city")).lower()
    state = text(row.get("state")).lower()
    return (any(s in state for s in CORE_STATES)
            or any(c in city for c in CORE_CITIES))


def email_points(row, direct, generic):
    email = row["_email"]
    if not email:
        return 0
    points = generic if email.startswith(GENERIC_EMAILS) else direct
    if row["_verify"]:
        points = points // 2
    return points


def norm_name(value):
    words = re.findall(r"[a-z0-9]+", text(value).lower())
    return " ".join(w for w in words if w not in STOP_WORDS)


def company_type(row):
    raw = text(row.get("organiser_type")).lower().replace("_", " ")
    name = text(row.get("name")).lower()
    summary = text(row.get("summary")).lower()

    # Venue by name first, so UI Conference Centre stays a venue
    if any(w in name for w in VENUE_NAME_WORDS):
        return "Venue"

    # Schools, churches, banks, government, political parties
    if any(w in name or w in raw for w in EXCLUDE_WORDS):
        return "Other"

    if "venue" in raw:
        return "Venue"

    # Read type, name AND summary together
    blob = f"{raw} {name} {summary}"

    if any(w in blob for w in ("production", "rental", "staging", "audio visual")):
        return "Production company"
    if any(w in blob for w in ("agency", "experiential", "activation",
                               "marketing", "brand", "communication")):
        return "Agency"
    if any(w in blob for w in ("event", "planner", "organiser", "organizer",
                               "exhibition", "chamber", "concept")):
        return "Event organiser"

    return "Other"


def priority(score):
    if score >= 80:
        return "Hot"
    if score >= 60:
        return "Strong"
    if score >= 40:
        return "Research"
    return "Low"


# ---------- scoring ----------

def score_event(row):
    score = 0
    start = parse_date(row.get("start_date"))

    # Best window: 2 weeks to 4 months out
    if start:
        days = (start - TODAY).days
        if days <= 13:
            score += 10
        elif days <= 120:
            score += 25
        elif days <= 240:
            score += 15
        else:
            score += 8

    score += email_points(row, 12, 6)
    if row["_phone"]:
        score += 10

    score += {"large": 15, "medium": 8}.get(text(row.get("scale")).lower(), 0)

    needs = len(row.get("needs") or [])
    score += 15 if needs >= 4 else 10 if needs >= 2 else 5 if needs == 1 else 0

    score += 10 if in_core_area(row) else -5

    otype = text(row.get("organiser_type")).lower()
    if otype in ("agency", "production company", "brand"):
        score += 10
    elif otype in ("government", "church", "university", "venue"):
        score += 5

    confidence = text(row.get("confidence")).lower()
    score += 5 if confidence == "high" else -5 if confidence == "low" else 0

    if len(row["_links"]) >= 2:
        score += 5

    score = max(0, min(100, score))

    # No confirmed date, never above Research
    if not start:
        score = min(score, 59)

    return score


TYPE_POINTS = {"Agency": 30, "Production company": 30, "Venue": 25,
               "Event organiser": 22, "Other": 0}


def score_company(row):
    score = TYPE_POINTS[row["_type"]]

    score += email_points(row, 20, 12)
    if row["_phone"]:
        score += 15
    if in_core_area(row):
        score += 15
    if text(row.get("website")):
        score += 10

    confidence = text(row.get("confidence")).lower()
    score += 10 if confidence == "high" else -5 if confidence == "low" else 0

    score = max(0, min(100, score))

    # Schools, banks, political parties and similar are not the target
    if row["_type"] == "Other":
        score = min(score, 39)

    return score


# ---------- deduplication ----------

def same_event(a, b):
    na, nb = a["_norm"], b["_norm"]
    if not na or not nb:
        return False

    ratio = SequenceMatcher(None, na, nb).ratio()
    contained = min(len(na), len(nb)) >= 8 and (na in nb or nb in na)
    if ratio < 0.8 and not contained:
        return False

    da, db = parse_date(a.get("start_date")), parse_date(b.get("start_date"))
    if da and db and abs((da - db).days) > 10:
        return False

    ca, cb = text(a.get("city")).lower(), text(b.get("city")).lower()
    if ca and cb and ca not in cb and cb not in ca:
        return False

    return True


def same_company(a, b):
    da, db = domain(a.get("website")), domain(b.get("website"))
    if da and db and da == db:
        return True
    na, nb = a["_norm"], b["_norm"]
    return bool(na and nb and SequenceMatcher(None, na, nb).ratio() >= 0.85)


def completeness(row):
    filled = sum(1 for f in MAIN_FIELDS if text(row.get(f)))
    return filled + (1 if row["_email"] else 0) + (1 if row["_phone"] else 0)


def merge(group):
    group.sort(key=completeness, reverse=True)
    main = dict(group[0])

    for other in group[1:]:
        for field in MAIN_FIELDS + ["event_type", "state", "organiser_type",
                                    "summary", "scale"]:
            if not text(main.get(field)) and text(other.get(field)):
                main[field] = other[field]
        if not main["_email"] and other["_email"]:
            main["_email"] = other["_email"]
        if not main["_phone"] and other["_phone"]:
            main["_phone"] = other["_phone"]
        main["needs"] = sorted(set(main.get("needs") or [])
                               | set(other.get("needs") or []))

    links = []
    for row in group:
        for link in row["_links"]:
            if link not in links:
                links.append(link)
    main["_links"] = links
    return main


def dedupe(rows, same):
    groups = []
    for row in rows:
        for group in groups:
            if same(group[0], row):
                group.append(row)
                break
        else:
            groups.append([row])
    return [merge(g) for g in groups]


# ---------- excel ----------

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF")
PRIORITY_FILLS = {"Hot": "FECACA", "Strong": "FED7AA",
                  "Research": "FEF08A", "Low": "E5E7EB"}
VERIFY_FILL = PatternFill("solid", fgColor="FEF3C7")

EVENT_HEADERS = ["Priority", "Score", "Event", "Type", "Date", "Days Away",
                 "Venue", "City", "State", "Organiser", "Organiser Type",
                 "Email", "Phone", "Verify Contact", "Website",
                 "Likely Needs", "Scale", "Summary", "Sources", "Link",
                 "Status", "Notes"]
EVENT_WIDTHS = [10, 7, 38, 13, 13, 10, 30, 14, 10, 28, 16, 30, 14, 12,
                28, 30, 9, 50, 8, 40, 14, 25]

COMPANY_HEADERS = ["Priority", "Score", "Company", "Type", "City", "State",
                   "Email", "Phone", "Verify Contact", "Website",
                   "Summary", "Link", "Status", "Notes"]
COMPANY_WIDTHS = [10, 7, 34, 18, 14, 10, 30, 14, 12, 30, 50, 40, 14, 25]


def write_sheet(ws, headers, rows, widths, link_col, wrap_col, verify_col):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row in rows:
        ws.append(row)

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    for row in ws.iter_rows(min_row=2):
        label = row[0].value
        if label in PRIORITY_FILLS:
            row[0].fill = PatternFill("solid", fgColor=PRIORITY_FILLS[label])
            row[0].font = Font(bold=True)

        for cell in row:
            cell.alignment = Alignment(vertical="top")

        row[wrap_col - 1].alignment = Alignment(vertical="top", wrap_text=True)

        if row[verify_col - 1].value:
            row[verify_col - 1].fill = VERIFY_FILL

        link_cell = row[link_col - 1]
        if link_cell.value:
            link_cell.hyperlink = link_cell.value
            link_cell.font = Font(color="2563EB", underline="single")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def event_row(r):
    start = parse_date(r.get("start_date"))
    return [
        priority(r["_score"]), r["_score"], text(r.get("name")),
        text(r.get("event_type")),
        start.isoformat() if start else "Not confirmed",
        (start - TODAY).days if start else "",
        text(r.get("venue")), text(r.get("city")), text(r.get("state")),
        text(r.get("organiser")), text(r.get("organiser_type")),
        r["_email"], r["_phone"], r["_verify"], text(r.get("website")),
        ", ".join(r.get("needs") or []), text(r.get("scale")),
        text(r.get("summary")), len(r["_links"]), r["_links"][0],
        "", "",
    ]


def company_row(r):
    return [
        priority(r["_score"]), r["_score"], text(r.get("name")),
        r["_type"], text(r.get("city")), text(r.get("state")),
        r["_email"], r["_phone"], r["_verify"], text(r.get("website")),
        text(r.get("summary")), r["_links"][0], "", "",
    ]


# ---------- main ----------

def main():
    if not os.path.exists(IN_PATH):
        print("No extracted.json found. Run extract.py first.")
        return

    with open(IN_PATH, "r", encoding="utf-8") as handle:
        rows = json.load(handle)

    events, companies = [], []
    too_soon = 0

    for row in rows:
        row["_email"] = pick_email(row)
        row["_phone"] = pick_phone(row)
        row["_links"] = [row["link"]]
        row["_norm"] = norm_name(row.get("name"))

        if row.get("lead_type") == "event":
            start = parse_date(row.get("start_date"))
            if start and (start - TODAY).days < MIN_DAYS_AWAY:
                too_soon += 1
                continue
            events.append(row)
        elif row.get("lead_type") == "company":
            companies.append(row)

    raw_events, raw_companies = len(events), len(companies)
    events = dedupe(events, same_event)
    companies = dedupe(companies, same_company)

    for row in events:
        row["_verify"] = needs_verify(row)
        row["_score"] = score_event(row)
    for row in companies:
        row["_verify"] = needs_verify(row)
        row["_type"] = company_type(row)
        row["_score"] = score_company(row)

    events.sort(key=lambda r: (-r["_score"],
                               parse_date(r.get("start_date")) or date.max))
    companies.sort(key=lambda r: -r["_score"])

    dated = [r for r in events if parse_date(r.get("start_date"))]
    undated = [r for r in events if not parse_date(r.get("start_date"))]
    good_companies = [r for r in companies if r["_type"] != "Other"]

    # ---- workbook ----
    wb = Workbook()

    ws = wb.active
    ws.title = "Event Leads"
    write_sheet(ws, EVENT_HEADERS, [event_row(r) for r in dated],
                EVENT_WIDTHS, link_col=20, wrap_col=18, verify_col=14)

    ws = wb.create_sheet("Company Leads")
    write_sheet(ws, COMPANY_HEADERS, [company_row(r) for r in companies],
                COMPANY_WIDTHS, link_col=12, wrap_col=11, verify_col=9)

    ws = wb.create_sheet("Date Not Confirmed")
    write_sheet(ws, EVENT_HEADERS, [event_row(r) for r in undated],
                EVENT_WIDTHS, link_col=20, wrap_col=18, verify_col=14)

    with_contact = sum(1 for r in dated if r["_email"] or r["_phone"])
    next_30 = sum(1 for r in dated
                  if (parse_date(r.get("start_date")) - TODAY).days <= 30)

    ws = wb.create_sheet("Summary")
    summary = [
        ("Generated", TODAY.isoformat()),
        ("", ""),
        ("Event leads (confirmed date)", len(dated)),
        ("  Hot", sum(1 for r in dated if r["_score"] >= 80)),
        ("  Strong", sum(1 for r in dated if 60 <= r["_score"] < 80)),
        ("  Happening in next 30 days", next_30),
        ("  With email or phone", with_contact),
        ("", ""),
        ("Events with date not confirmed", len(undated)),
        ("", ""),
        ("Company leads", len(companies)),
        ("  Agencies, production, venues, organisers", len(good_companies)),
        ("", ""),
        ("Housekeeping", ""),
        ("  Event duplicates merged", raw_events - len(events)),
        ("  Company duplicates merged", raw_companies - len(companies)),
        ("  Dropped, under 5 days away", too_soon),
        ("  Items reviewed by AI", len(rows)),
    ]
    for label, value in summary:
        ws.append([label, value])
        if label and not label.startswith(" "):
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14

    # Open on the Summary sheet
    wb.active = wb.sheetnames.index("Summary")

    os.makedirs(OUT_DIR, exist_ok=True)
    outfile = os.path.join(OUT_DIR, f"leads_{TODAY.isoformat()}.xlsx")

    try:
        wb.save(outfile)
    except PermissionError:
        print(f"Could not save. Close {outfile} in Excel and run again.")
        return

    print(f"Event leads:     {len(dated)} with confirmed date")
    print(f"Date unclear:    {len(undated)} (separate sheet)")
    print(f"Company leads:   {len(companies)} ({len(good_companies)} real targets)")
    print(f"With contact:    {with_contact} of {len(dated)} dated events")
    print()
    print("TOP 10 EVENT LEADS")
    print("-" * 60)
    for r in dated[:10]:
        print(f"  {r['_score']:>3}  {text(r.get('start_date')):<10}  "
              f"{text(r.get('name'))[:40]}")
    print()
    print("TOP 10 COMPANY LEADS")
    print("-" * 60)
    for r in companies[:10]:
        print(f"  {r['_score']:>3}  {r['_type']:<18}  {text(r.get('name'))[:35]}")
    print()
    print(f"Saved to: {outfile}")


if __name__ == "__main__":
    main()